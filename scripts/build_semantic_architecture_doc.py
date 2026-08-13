#!/usr/bin/env python3
"""Build the reviewed OFL semantic-entity architecture as a polished DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "OPEN_FORECAST_LIBRARY_SEMANTIC_ENTITY_ARCHITECTURE_V0_1.md"
OUTPUT = ROOT / "docs" / "OPEN_FORECAST_LIBRARY_SEMANTIC_ENTITY_ARCHITECTURE_V0_1.docx"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

INK = "0F172A"
BODY = "263244"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "64748B"
LIGHT_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
GREEN_FILL = "ECFDF5"
GREEN = "047857"
GOLD_FILL = "FFFBEB"
GOLD = "92400E"
BORDER = "CBD5E1"


def add_abstract_numbering(doc, *, kind):
    """Create a real Word numbering definition for bullets or decimals."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    return abstract_id


def add_numbering_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    # Explicitly restart each distinct Markdown list at 1. LibreOffice may
    # otherwise continue numbering across separate numId instances that share
    # one abstract definition.
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    num.append(lvl_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)


def finish_list_paragraph(paragraph):
    """Keep list items visibly separate in both Word and LibreOffice renders."""
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.10
    # LibreOffice can visually run consecutive numbered paragraphs together
    # when custom numbering definitions are used. A final soft break makes the
    # boundary explicit without changing the list semantics in the DOCX.
    paragraph.add_run().add_break()


def add_transition_spacer(doc):
    """Insert a compact paragraph to prevent renderer run-in text."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = Pt(8)
    run = p.add_run(" ")
    set_run_font(run, size=1, color="FFFFFF")


def set_run_font(run, *, name="Arial", size=11, color=BODY, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color=BORDER, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def column_widths(rows):
    count = len(rows[0])
    if count == 2:
        return [2600, 6760]
    if count == 3:
        return [2200, 2700, 4460]
    if count == 4:
        return [1850, 2400, 2600, 2510]
    weights = []
    for col in range(count):
        maximum = max(len(row[col]) for row in rows)
        weights.append(max(8, min(maximum, 42)))
    total = sum(weights)
    widths = [max(1100, int(CONTENT_WIDTH_DXA * weight / total)) for weight in weights]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("OPEN FORECAST LIBRARY  |  ARCHITECTURE PROPOSAL"), size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    set_run_font(p.add_run("Future Edge Group FZE  |  "), size=8.5, color=MUTED)
    add_page_number(p)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 12, 6),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
        ("Heading 4", 11.5, DARK_BLUE, 7, 3),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.05

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://\S+)")


def add_inline(paragraph, text, *, base_size=11, base_color=BODY):
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            set_run_font(paragraph.add_run(text[cursor:match.start()]), size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), size=base_size, color=base_color, bold=True)
        elif token.startswith("`"):
            set_run_font(paragraph.add_run(token[1:-1]), name="Consolas", size=max(8.5, base_size - 1), color=DARK_BLUE)
        else:
            clean = token.rstrip(".,)")
            trailing = token[len(clean):]
            set_run_font(paragraph.add_run(clean), size=max(8.5, base_size - 1), color=BLUE)
            if trailing:
                set_run_font(paragraph.add_run(trailing), size=base_size, color=base_color)
        cursor = match.end()
    if cursor < len(text):
        set_run_font(paragraph.add_run(text[cursor:]), size=base_size, color=base_color)


def add_callout(doc, text, fill=GREEN_FILL, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=accent, size="8")
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text, base_color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=BORDER, size="4")
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    set_cell_shading(cell, "F8FAFC")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("\n".join(lines))
    set_run_font(run, name="Consolas", size=8.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    set_table_geometry(table, column_widths(rows))
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, value, base_size=9.2, base_color=INK)
            if row_index == 0:
                for run in p.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def parse_markdown_body(doc, lines):
    index = 0
    in_code = False
    code_lines = []
    bullet_abstract_id = add_abstract_numbering(doc, kind="bullet")
    decimal_abstract_id = add_abstract_numbering(doc, kind="decimal")
    current_list_kind = None
    current_num_id = None
    in_sources = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            if current_list_kind is not None:
                add_transition_spacer(doc)
            current_list_kind = None
            current_num_id = None
            index += 1
            continue
        if stripped.startswith("| ") and index + 1 < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[index + 1].strip()):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                raw = lines[index].strip()
                if not re.match(r"^\|[-:| ]+\|$", raw):
                    table_rows.append([cell.strip() for cell in raw.strip("|").split("|")])
                index += 1
            add_markdown_table(doc, table_rows)
            continue
        if stripped.startswith("### "):
            current_list_kind = None
            current_num_id = None
            add_transition_spacer(doc)
            doc.add_paragraph(stripped[4:], style="Heading 2")
        elif stripped.startswith("#### "):
            current_list_kind = None
            current_num_id = None
            add_transition_spacer(doc)
            doc.add_paragraph(stripped[5:], style="Heading 3")
        elif stripped.startswith("## "):
            current_list_kind = None
            current_num_id = None
            add_transition_spacer(doc)
            doc.add_paragraph(stripped[3:], style="Heading 1")
            in_sources = stripped == "## 14. Sources"
        elif stripped.startswith("# "):
            current_list_kind = None
            current_num_id = None
            pass
        elif stripped.startswith("> "):
            current_list_kind = None
            current_num_id = None
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            add_callout(doc, "  |  ".join(quote_lines))
            continue
        elif re.match(r"^- ", stripped):
            if current_list_kind != "bullet":
                current_list_kind = "bullet"
                current_num_id = add_numbering_instance(doc, bullet_abstract_id)
            p = doc.add_paragraph()
            apply_numbering(p, current_num_id)
            add_inline(p, stripped[2:])
            finish_list_paragraph(p)
        elif re.match(r"^\d+\. ", stripped):
            if current_list_kind != "decimal":
                current_list_kind = "decimal"
                current_num_id = add_numbering_instance(doc, decimal_abstract_id)
            p = doc.add_paragraph()
            apply_numbering(p, current_num_id)
            add_inline(p, re.sub(r"^\d+\. ", "", stripped), base_size=9.2 if in_sources else 11)
            if in_sources:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
            else:
                finish_list_paragraph(p)
        elif stripped == "---":
            index += 1
            continue
        else:
            current_list_kind = None
            current_num_id = None
            p = doc.add_paragraph()
            add_inline(p, stripped)
        index += 1


def add_front_matter(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("ARCHITECTURE DECISION MEMO"), size=9, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run("Open Forecast Library"), size=26, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    set_run_font(p.add_run("Semantic entity, publisher, database, and blockchain proof architecture"), size=14, color=MUTED)

    metadata = [
        ("Status", "Implemented staging baseline"),
        ("Version", "0.1"),
        ("Date", "12 August 2026"),
        ("Prepared for", "Future Edge Group FZE / iPulse AI"),
        ("Decision", "Adopt Entity -> Target -> Forecast -> Receipt -> Proof -> Evaluation"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)
    set_table_geometry(table, [1600, 7760])
    set_table_borders(table, color="E2E8F0", size="2")
    for row, (label, value) in zip(table.rows, metadata):
        set_cell_shading(row.cells[0], LIGHT_FILL)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=65, start=110, bottom=65, end=110)
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(label), size=9, color=DARK_BLUE, bold=True)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size=9.2, color=INK)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_callout(
        doc,
        "Recommendation: keep Firestore as the governed operational catalog, use BigQuery only as an analytics mirror, and publish blockchain proofs through a dedicated KMS-controlled Base signer. iPulse AI is the first publisher profile, not the name of the Library.",
        fill=GREEN_FILL,
        accent=GREEN,
    )


def build():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    configure_section(doc.sections[0])
    configure_styles(doc)
    add_front_matter(doc)

    body_start = next(i for i, line in enumerate(lines) if line.startswith("## Executive decision"))
    parse_markdown_body(doc, lines[body_start:])

    core = doc.core_properties
    core.title = "Open Forecast Library Semantic Entity Architecture"
    core.subject = "Entity, publisher, Firestore, BigQuery, and Base/EAS architecture proposal"
    core.author = "Future Edge Group FZE"
    core.keywords = "Open Forecast Library, Open Forecast Receipt, iPulse AI, semantic web, Firestore, Wikidata, EAS, Base"
    core.comments = "Reviewed architecture and implemented staging semantic-entity foundation."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
